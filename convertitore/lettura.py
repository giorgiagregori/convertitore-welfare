"""Lettura del file del provider (xlsx, xls o csv) in una lista di righe.

Il risultato è sempre lo stesso, qualunque sia il formato:
    una lista di RigaFile(numero, valori)
dove `numero` è il numero di riga nel file originale (come lo vede chi apre il
file in Excel) e `valori` è la lista dei valori delle celle.

Qui si riconoscono da soli:
    - il tipo di file (dai primi byte, non solo dall'estensione)
    - per i csv: la codifica (UTF-8 o Windows) e il separatore (; , tab |)
Le righe completamente vuote vengono saltate, ma la numerazione resta quella
originale, così negli scarti "riga 8" è davvero la riga 8 del file.
"""

import csv
import io
from dataclasses import dataclass


@dataclass
class RigaFile:
    numero: int
    valori: list


class FileNonLeggibile(ValueError):
    """Il file non si può leggere: formato sconosciuto o rovinato."""


def tipo_file(contenuto: bytes, nome_file: str = "") -> str:
    """Riconosce il tipo dai primi byte: un file .xlsx è uno zip, un .xls è un
    contenitore OLE di Microsoft; tutto il resto lo trattiamo come testo (csv)."""
    if contenuto[:4] == b"PK\x03\x04":
        return "xlsx"
    if contenuto[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    estensione = nome_file.lower().rsplit(".", 1)[-1] if "." in nome_file else ""
    if estensione in ("xlsx", "xls"):
        raise FileNonLeggibile(f"Il file '{nome_file}' non è un Excel valido")
    return "csv"


def leggi_righe(contenuto: bytes, nome_file: str = "") -> list:
    tipo = tipo_file(contenuto, nome_file)
    if tipo == "xlsx":
        righe = _leggi_xlsx(contenuto)
    elif tipo == "xls":
        righe = _leggi_xls(contenuto)
    else:
        righe = _leggi_csv(contenuto)
    return [r for r in righe if any(_pieno(v) for v in r.valori)]


def _pieno(valore) -> bool:
    return valore is not None and str(valore).strip() != ""


def _leggi_xlsx(contenuto: bytes) -> list:
    from openpyxl import load_workbook
    try:
        libro = load_workbook(io.BytesIO(contenuto), read_only=True, data_only=True)
    except Exception as errore:
        raise FileNonLeggibile(f"Excel non leggibile: {errore}")
    foglio = libro.worksheets[0]
    righe = [RigaFile(i, list(valori)) for i, valori in enumerate(foglio.iter_rows(values_only=True), start=1)]
    libro.close()
    return righe


def _leggi_xls(contenuto: bytes) -> list:
    try:
        import xlrd
    except ImportError:
        raise FileNonLeggibile("Per leggere i file .xls serve la libreria xlrd (vedi requirements.txt)")
    try:
        libro = xlrd.open_workbook(file_contents=contenuto)
    except Exception as errore:
        raise FileNonLeggibile(f"Excel .xls non leggibile: {errore}")
    foglio = libro.sheet_by_index(0)
    righe = []
    for i in range(foglio.nrows):
        valori = []
        for cella in foglio.row(i):
            valore = cella.value
            # xlrd restituisce i numeri interi come float (20 -> 20.0)
            if cella.ctype == xlrd.XL_CELL_NUMBER and float(valore).is_integer():
                valore = int(valore)
            valori.append(valore)
        righe.append(RigaFile(i + 1, valori))
    return righe


def _decodifica(contenuto: bytes) -> str:
    """UTF-8 (con o senza BOM) se possibile, altrimenti la codifica di Windows
    (cp1252), tipica dei file esportati da gestionali italiani: è il caso del
    Provider B, dove "Liberalità" non è UTF-8."""
    try:
        return contenuto.decode("utf-8-sig")
    except UnicodeDecodeError:
        return contenuto.decode("cp1252", errors="replace")


def _leggi_csv(contenuto: bytes) -> list:
    testo = _decodifica(contenuto)
    separatore = _separatore(testo)
    lettore = csv.reader(io.StringIO(testo), delimiter=separatore)
    return [RigaFile(i, [v.strip() for v in valori]) for i, valori in enumerate(lettore, start=1)]


def _separatore(testo: str) -> str:
    """Sceglie il separatore dei csv.

    Non uso csv.Sniffer: con importi come "60,00" prende la virgola dei
    decimali per separatore (file del Provider B). Tra i separatori costanti su
    almeno l'80% delle righe vince quello con più colonne; a parità, ";".
    """
    righe = [r for r in testo.splitlines()[:30] if r.strip()]
    migliore, punteggio_migliore = ";", (False, 0, 0)
    for separatore in (";", "\t", "|", ","):
        colonne = [len(next(csv.reader([r], delimiter=separatore))) for r in righe]
        if not colonne or max(colonne) < 2:
            continue
        piu_frequente = max(set(colonne), key=colonne.count)
        costanza = colonne.count(piu_frequente) / len(colonne)
        punteggio = (costanza >= 0.8, piu_frequente, costanza)
        if punteggio > punteggio_migliore:
            migliore, punteggio_migliore = separatore, punteggio
    return migliore
